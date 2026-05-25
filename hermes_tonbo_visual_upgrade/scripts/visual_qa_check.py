from pathlib import Path
from pptx import Presentation
import json

pptx_path = Path(r"C:/Users/gagan/Desktop/hermes+fabric/hermes_tonbo_visual_upgrade/output/Tonbo_Investor_DD_Deck_visual_v3.pptx")
prs = Presentation(str(pptx_path))

report=[]
flags=[]
for i,slide in enumerate(prs.slides, start=1):
    txt_chars=0
    image_count=0
    chart_count=0
    shape_count=len(slide.shapes)
    for sh in slide.shapes:
        if getattr(sh,'has_text_frame',False) and sh.has_text_frame:
            txt_chars += sum(len(p.text or '') for p in sh.text_frame.paragraphs)
        if sh.shape_type == 13:  # picture
            image_count += 1
        if getattr(sh,'has_chart',False) and sh.has_chart:
            chart_count += 1
    has_visual = (image_count + chart_count) > 0
    memo_like = txt_chars > 850 and (image_count + chart_count) < 2
    if not has_visual:
        flags.append({'slide':i,'issue':'no_visual'})
    if memo_like:
        flags.append({'slide':i,'issue':'text_heavy'})
    report.append({'slide':i,'shape_count':shape_count,'text_chars':txt_chars,'images':image_count,'charts':chart_count,'has_visual':has_visual})

out = {
    'deck': str(pptx_path),
    'slides': len(prs.slides),
    'flags': flags,
    'per_slide': report
}
json_path = pptx_path.parent / 'visual_qa_report.json'
json_path.write_text(json.dumps(out, indent=2), encoding='utf-8')
print(json_path)
print('flags', len(flags))
