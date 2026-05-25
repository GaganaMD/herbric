from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

BASE = Path(r"C:/Users/gagan/Desktop/hermes+fabric/hermes_tonbo_visual_upgrade")
OUT = BASE / "output"
CH = OUT / "charts"
SEL = BASE / "assets_selected"
OUT.mkdir(parents=True, exist_ok=True)

# Inputs from prior run
rev_chart = CH / "rev.png"
qoe_chart = CH / "qoe.png"
conc_chart = CH / "conc.png"
wc_chart = CH / "wc.png"
imgs = sorted([p for p in SEL.glob('anchor_*')])

def pic(i):
    return str(imgs[i % len(imgs)]) if imgs else None

prs = Presentation()
# Force 16:9 widescreen canvas for institutional presentation mode
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W, H = prs.slide_width, prs.slide_height
# 16:9 widescreen inches
M_L, M_R, M_T, M_B = 0.45, 0.45, 0.32, 0.30
CONTENT_W = 13.33 - M_L - M_R
CONTENT_H = 7.5 - M_T - M_B

C_TITLE = RGBColor(15, 32, 60)
C_TEXT = RGBColor(35, 35, 35)
C_MUTE = RGBColor(90, 100, 112)
C_PANEL = RGBColor(246, 248, 251)
C_BORDER = RGBColor(215, 222, 232)


def tbox(slide, l, t, w, h, text, size=14, bold=False, color=C_TEXT, align=PP_ALIGN.LEFT):
    sh = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return sh


def title(slide, text):
    return tbox(slide, M_L, M_T, CONTENT_W, 0.62, text, size=28, bold=True, color=C_TITLE)


def panel(slide, l, t, w, h, header, bullets):
    sh = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = C_PANEL
    sh.line.color.rgb = C_BORDER
    tf = sh.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(0.14); tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.08); tf.margin_bottom = Inches(0.06)
    p = tf.paragraphs[0]
    p.text = header
    p.font.bold = True; p.font.size = Pt(15); p.font.color.rgb = C_TITLE
    for b in bullets:
        q = tf.add_paragraph(); q.text = f"• {b}"; q.font.size = Pt(12.5); q.font.color.rgb = C_TEXT
    return sh

# Slide 1 cover: contained image + dark ribbon for guaranteed contrast
s = prs.slides.add_slide(prs.slide_layouts[6])
if pic(0):
    s.shapes.add_picture(pic(0), Inches(M_L), Inches(M_T+0.15), width=Inches(CONTENT_W), height=Inches(5.9))
bar = s.shapes.add_shape(1, Inches(M_L), Inches(M_T+5.35), Inches(CONTENT_W), Inches(1.15))
bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor(9, 18, 36); bar.fill.transparency = 12; bar.line.fill.background()
tbox(s, M_L+0.25, M_T+5.50, CONTENT_W-0.5, 0.55, 'Tonbo Imaging | Investor Diligence Underwriting Deck', size=25, bold=True, color=RGBColor(255,255,255))
tbox(s, M_L+0.25, M_T+6.00, CONTENT_W-0.5, 0.35, 'Presentation-system refinement pass: readability, containment, executive mode', size=12, color=RGBColor(220,228,240))

# Slide 2 exec dashboard
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, 'Executive underwriting dashboard')
s.shapes.add_picture(str(rev_chart), Inches(M_L), Inches(1.10), width=Inches(6.0), height=Inches(2.55))
s.shapes.add_picture(str(conc_chart), Inches(6.85), Inches(1.10), width=Inches(6.0), height=Inches(2.55))
panel(s, M_L, 3.90, CONTENT_W, 2.70, 'IC bottom line', [
    'Growth is real; adjusted earnings remain negative after normalization.',
    'Concentration stack remains too high for clean underwriting sign-off.',
    'Proceed only with milestone-gated capital and control remediation.'
])

# Slide 3 operational context
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, 'Operational context: deployment reality and system footprint')
if pic(1): s.shapes.add_picture(pic(1), Inches(M_L), Inches(1.15), width=Inches(8.0), height=Inches(5.95))
panel(s, 8.65, 1.15, 4.23, 5.95, 'Visual interpretation', [
    'Read as a deployment-system business, not a pure software model.',
    'Execution risk links to hardware lifecycle and procurement cadence.',
    'Field reliability and integration depth drive investor confidence.'
])

# Slide 4 QoE
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, 'Quality of earnings: reported vs investable economics')
s.shapes.add_picture(str(qoe_chart), Inches(M_L), Inches(1.15), width=Inches(8.1), height=Inches(5.9))
panel(s, 8.70, 1.15, 4.18, 5.9, 'Underwriting consequence', [
    'FY25 EBITDA sign flips after normalization adjustments.',
    'Reported EBITDA multiple is not decision-grade.',
    'Require quarter-by-quarter path-to-profitability evidence.'
])

# Slide 5 concentration matrix
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, 'Concentration risk matrix (customer + SKU + vendor)')
box = s.shapes.add_shape(1, Inches(M_L), Inches(1.25), Inches(7.65), Inches(5.8))
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(253,253,253); box.line.color.rgb = C_BORDER
v = s.shapes.add_shape(1, Inches(4.25), Inches(1.25), Inches(0.02), Inches(5.8)); v.fill.solid(); v.fill.fore_color.rgb = RGBColor(190,190,190); v.line.fill.background()
h = s.shapes.add_shape(1, Inches(M_L), Inches(4.15), Inches(7.65), Inches(0.02)); h.fill.solid(); h.fill.fore_color.rgb = RGBColor(190,190,190); h.line.fill.background()
for (x,y,t,c) in [(0.95,1.75,'ALT Group\n49% rev',RGBColor(192,0,0)),(4.45,1.75,'MN105-AL\n68% SKU',RGBColor(237,125,49)),(0.95,4.55,'Branch/Vendor\ndependency',RGBColor(255,192,0)),(4.45,4.55,'Correlated\ndownside risk',RGBColor(112,173,71))]:
    tbox(s, x, y, 2.8, 1.4, t, size=16, bold=True, color=c)
if pic(2): s.shapes.add_picture(pic(2), Inches(8.25), Inches(1.25), width=Inches(4.63), height=Inches(3.45))
panel(s, 8.25, 4.95, 4.63, 2.1, 'Risk stance', ['Treat concentration as structural, not temporary variance.'])

# Slide 6 WC
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, 'Working-capital mechanics and liquidity fragility')
s.shapes.add_picture(str(wc_chart), Inches(M_L), Inches(1.15), width=Inches(6.35), height=Inches(3.10))
if pic(3): s.shapes.add_picture(pic(3), Inches(7.00), Inches(1.15), width=Inches(5.88), height=Inches(3.10))
panel(s, M_L, 4.45, CONTENT_W, 2.55, 'Interpretation', [
    'Current CCC is DPO-supported; normalize payables in downside case.',
    'Liquidity runway can compress quickly under vendor-term reset.'
])

# Slide 7 governance heatmap
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, 'Governance/control heatmap')
rows=[('Accounting policy','High'),('Inventory control','High'),('Tax/GST hygiene','High'),('Close discipline','Medium'),('Forecast reliability','Medium')]
for i,(a,b) in enumerate(rows):
    y = 1.20 + i*1.05
    sh=s.shapes.add_shape(1, Inches(M_L), Inches(y), Inches(7.55), Inches(0.82))
    sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor(246,248,251); sh.line.color.rgb=C_BORDER
    tbox(s, M_L+0.15, y+0.16, 4.6, 0.42, a, size=13)
    risk=RGBColor(192,0,0) if b=='High' else RGBColor(237,125,49)
    tag=s.shapes.add_shape(1, Inches(6.15), Inches(y+0.13), Inches(1.25), Inches(0.56))
    tag.fill.solid(); tag.fill.fore_color.rgb=risk; tag.line.fill.background()
    tbox(s, 6.16, y+0.19, 1.2, 0.30, b, size=11.5, bold=True, color=RGBColor(255,255,255), align=PP_ALIGN.CENTER)
if pic(4): s.shapes.add_picture(pic(4), Inches(8.05), Inches(1.20), width=Inches(4.83), height=Inches(5.02))

# Slide 8 scenarios
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, 'Scenario underwriting: base / downside / severe')
cards=[('Base',RGBColor(112,173,71),'Moderate growth with partial control fixes.'),('Downside',RGBColor(237,125,49),'DPO normalization and slower collections.'),('Severe',RGBColor(192,0,0),'Concentration shock and compliance cash outflow.')]
for i,(name,col,txt) in enumerate(cards):
    x=M_L+i*4.25
    sh=s.shapes.add_shape(1, Inches(x), Inches(1.65), Inches(4.0), Inches(4.25))
    sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor(251,251,252); sh.line.color.rgb=col
    tbox(s, x+0.20, 1.92, 3.6, 0.5, name, size=20, bold=True, color=col)
    tbox(s, x+0.20, 2.55, 3.55, 2.8, txt, size=12.5)
if pic(5): s.shapes.add_picture(pic(5), Inches(M_L), Inches(6.05), width=Inches(CONTENT_W), height=Inches(0.95))

# Slide 9 capability stack
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, 'Capability stack and ecosystem positioning')
labels=[('Sensor\nHardware',M_L),('Embedded\nCompute',3.55),('Analytics\nControl',6.20),('Deployment\nPlatform',8.85)]
for t,x in labels:
    sh=s.shapes.add_shape(1, Inches(x), Inches(2.45), Inches(2.25), Inches(1.75))
    sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor(31,78,121); sh.line.fill.background()
    tbox(s, x+0.12, 2.95, 2.0, 0.85, t, size=13, bold=True, color=RGBColor(255,255,255), align=PP_ALIGN.CENTER)
if pic(6): s.shapes.add_picture(pic(6), Inches(M_L), Inches(1.10), width=Inches(3.15), height=Inches(1.15))
if pic(7): s.shapes.add_picture(pic(7), Inches(10.18), Inches(4.55), width=Inches(2.70), height=Inches(1.70))
panel(s, M_L, 4.70, 8.95, 2.3, 'Strategic read', ['Differentiation must convert into diversified revenue, not concentrated growth.'])

# Slide 10 IC decision panel
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, 'IC decision panel and conditions precedent')
panel(s, M_L, 1.2, 6.55, 5.85, 'Go / no-go logic', [
    'GO only with milestone-gated deployment and hard covenants.',
    'CPs: QoE closure, control remediation tracker, compliance closure plan.',
    'Valuation anchored to scenario-weighted outcomes, not headline EBITDA.'
])
if pic(8): s.shapes.add_picture(pic(8), Inches(7.15), Inches(1.2), width=Inches(5.73), height=Inches(5.85))

# Slide 11 100-day workflow
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, 'Post-close 100-day value-protection workflow')
if pic(9): s.shapes.add_picture(pic(9), Inches(M_L), Inches(1.10), width=Inches(CONTENT_W), height=Inches(1.35))
steps=['Monthly close hardening','Concentration PMO','Cash cockpit','Compliance closure office']
for i,st in enumerate(steps):
    x=M_L+i*3.15
    sh=s.shapes.add_shape(1, Inches(x), Inches(3.10), Inches(2.65), Inches(1.45))
    sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor(244,246,249); sh.line.color.rgb=C_BORDER
    tbox(s, x+0.12, 3.55, 2.4, 0.75, st, size=12, align=PP_ALIGN.CENTER)
    if i<3:
        ar=s.shapes.add_shape(1, Inches(x+2.65), Inches(3.78), Inches(0.40), Inches(0.05)); ar.fill.solid(); ar.fill.fore_color.rgb=RGBColor(110,110,110); ar.line.fill.background()
panel(s, M_L, 5.05, CONTENT_W, 1.95, 'Execution objective', ['Reduce reporting variance, concentration fragility, and liquidity opacity inside first 100 days.'])

# Slide 12 evidence
s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, 'Evidence boundaries and source traceability')
if pic(0): s.shapes.add_picture(pic(0), Inches(8.35), Inches(1.20), width=Inches(4.53), height=Inches(5.85))
panel(s, M_L, 1.20, 7.62, 5.85, 'Primary evidence', [
    'Workbook metrics + bridge combined context + Hermes synthesis output.',
    'Firecrawl retrieval and image-source manifests retained in task folder.',
    'Visuals are grounding aids; underwriting decisions rest on numeric evidence.',
    'Non-official imagery is tagged with source URL in selection manifest.'
])

out = OUT / 'Tonbo_Investor_DD_Deck_visual_v4_safe.pptx'
prs.save(out)
print(out)
